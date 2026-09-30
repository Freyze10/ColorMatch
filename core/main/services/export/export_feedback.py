import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse
from datetime import datetime, date
from main.utils.log_audit_trail import log_audit
from main.models import (
    tbl_feedback_details, tbl_cmf, tbl_cmf_dates, 
    tbl_cmf_formula, tbl_cmf_pending_completed, tbl_resins_selected,
    tbl_cmf_process02, tbl_cmf_color_req, tbl_mb_extruder_formula,
    tbl_dc_extruder_formula, tbl_submitted_selected
)


def val_or_default(val, is_date=False):
    """Returns '----' for any empty, null, or whitespace-only value."""
    if val is None:
        return "----"
    
    if is_date and isinstance(val, (datetime, date)):
        return val.strftime('%m/%d/%Y')
    
    str_val = str(val).strip()
    if not str_val or str_val.lower() == 'none':
        return "----"
    return str_val


def generate_feedback_excel(date_from=None, date_to=None):
    """
    Generates Excel for Feedback Records strictly for CMF.
    Exclusively uses tbl_cmf_dates for sorting and filtering.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feedback Records"

    # 1. Define Headers (Submitted Options placed immediately after AR #)
    headers = [
        "Matching No.", "Customer", "Date Created", "Date Lab Received", "Target Date",
        "Primary Color", "Color Description", "End Product", "Matching Type", "Salesman",
        "Color Req.", "Resin", "Process", "Type of Colorant", "Date Given Sample",
        "Set/PC", "Quantity Given", "Code Submitted", "Dosage", "Lot #", "AR #",
        "Submitted Options", "Status", "Date Standard & Result", "Comment", "Storage Details"
    ]

    # 2. Header Styles
    header_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
    header_font = Font(bold=True)
    center_align = Alignment(horizontal="center", vertical="center")

    for col_num, column_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=column_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # 3. Fetch Base Feedback Records (Strictly CMF)
    feedback_qs = (
        tbl_feedback_details.objects
        .filter(cm_no__isnull=False)
        .select_related('cm_no', 'cm_no__sm', 'cm_no__in_code_no')
        .all()
    )
    
    rows_to_sort = []

    for fb in feedback_qs:
        cmf = fb.cm_no
        if not cmf:
            continue

        # Single source of truth for dates
        dates_rec = tbl_cmf_dates.objects.filter(cm_no=cmf).first()
        created_date = dates_rec.form_made if dates_rec else None

        # Filter by Date Range if provided
        if date_from and date_to:
            if not created_date or not (date_from <= created_date <= date_to):
                continue

        # Pending / Completed Tracking
        pending = tbl_cmf_pending_completed.objects.filter(cm_no=cmf).select_related('code').first()

        # Submitted Options from tbl_submitted_selected
        submitted_options_list = list(
            tbl_submitted_selected.objects
            .filter(completed_id=pending)
            .values_list('option_id__name', flat=True)
        ) if pending else []
        submitted_options_str = ", ".join(submitted_options_list) if submitted_options_list else ""

        # Color Requirement
        color_req = tbl_cmf_color_req.objects.filter(cm_no=cmf).first()
        color_req_val = color_req.name if color_req else ""

        # Formula Specifications
        formula = tbl_cmf_formula.objects.filter(cm_no=cmf).first()
        customer = formula.customer if formula else ""
        end_product = formula.finished_product if formula else ""
        dosage = formula.dosage if formula else ""
        salesman = cmf.sm.name if cmf.sm else ""
        primary_color = cmf.in_code_no.color if cmf.in_code_no else ""

        # Process & Resins
        process_list = tbl_cmf_process02.objects.filter(cmf_formula_no=formula).values_list('process_no__name', flat=True) if formula else []
        processes = ", ".join(filter(None, process_list))

        resins_list = tbl_resins_selected.objects.filter(cm_no=cmf).values_list('resin_no__abbreviation', flat=True)
        resins = ", ".join(filter(None, resins_list))

        # Final Code & Lot Info (Priority: MB Final -> DC Final -> Tracking record)
        code_sub = ""
        lot_no = ""
        final_f = tbl_mb_extruder_formula.objects.filter(cm_no=cmf, is_final=True).select_related('code').first()
        
        if final_f:
            code_sub = final_f.code.product_code if final_f.code else ""
            lot_no = final_f.lot_no or ""
        else:
            final_dc = tbl_dc_extruder_formula.objects.filter(cm_no=cmf, is_final=True).select_related('code').first()
            if final_dc:
                code_sub = final_dc.code.product_code if final_dc.code else ""
            elif pending and pending.code:
                code_sub = pending.code.product_code
            
            if pending and pending.lot_no:
                lot_no = pending.lot_no

        # Format Dosage display
        dosage_str = f"{float(dosage):g}%" if dosage else ""

        # Build Data Row with default '----' for empty fields
        data_row = [
            val_or_default(cmf.cm_no),
            val_or_default(customer),
            val_or_default(created_date, is_date=True),
            val_or_default(dates_rec.date_received_lab if dates_rec else None),
            val_or_default(dates_rec.due_date_lab if dates_rec else None, is_date=True),
            val_or_default(primary_color),
            val_or_default(cmf.color_desc),
            val_or_default(end_product),
            val_or_default(cmf.matching_type),
            val_or_default(salesman),
            val_or_default(color_req_val),
            val_or_default(resins),
            val_or_default(processes),
            val_or_default(cmf.colorant_type),
            val_or_default(pending.date_submitted if pending else None, is_date=True),
            val_or_default(fb.pieces),
            val_or_default(fb.quantity_given),
            val_or_default(code_sub),
            val_or_default(dosage_str),
            val_or_default(lot_no),
            val_or_default(pending.ar_no if pending else None),
            val_or_default(submitted_options_str),
            val_or_default(fb.status),
            val_or_default(fb.date_sample_received, is_date=True),
            val_or_default(fb.comment),
            val_or_default(fb.storage_details)
        ]

        # Use 1900-01-01 as fallback for missing dates to keep them sorted predictably at top
        rows_to_sort.append((created_date or date(1900, 1, 1), data_row))

    # 4. Sort Ascending (Oldest at top, Latest at bottom)
    rows_to_sort.sort(key=lambda x: x[0])

    # 5. Write to Sheet
    for idx, (dt, row_data) in enumerate(rows_to_sort, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=idx, column=col_idx, value=value)
            if value == "----":
                cell.alignment = center_align

    # 6. Auto-fit Column Widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length: 
                    max_length = len(str(cell.value))
            except: 
                pass
        ws.column_dimensions[column].width = max(min(max_length + 3, 50), 12)

    return wb, rows_to_sort


def export_feedback_excel(request):
    from_date_str = request.GET.get('from')
    to_date_str = request.GET.get('to')
    
    date_from = None
    date_to = None
    
    try:
        if from_date_str: 
            date_from = datetime.strptime(from_date_str, '%m/%d/%Y').date()
        if to_date_str: 
            date_to = datetime.strptime(to_date_str, '%m/%d/%Y').date()
    except ValueError: 
        pass

    wb, sorted_data = generate_feedback_excel(date_from, date_to)
    
    # 7. Filename Construction (MMDDYY-MMDDYY)
    if date_from and date_to:
        f_range = f"{date_from:%m%d%y}-{date_to:%m%d%y}"
    elif sorted_data and sorted_data[0][0] != date(1900, 1, 1):
        first_date = sorted_data[0][0].strftime('%m%d%y')
        last_date = sorted_data[-1][0].strftime('%m%d%y')
        f_range = f"{first_date}-{last_date}"
    else:
        f_range = "All_Records"

    filename = f"Feedback_Report_{f_range}.xlsx"
    
    # 8. Log Audit Action
    range_desc = f"{from_date_str} to {to_date_str}" if from_date_str and to_date_str else "All Records"
    total_count = len(sorted_data)
    audit_details = f"Exported Feedback Records to Excel. Range: {range_desc}. Total: {total_count} records."

    log_audit(request, "Exported", audit_details)
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    
    return response